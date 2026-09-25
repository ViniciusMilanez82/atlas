"""Format-aware extraction with provenance (A3-09, spec 10.1-10.2; scenario T12).

Pure functions: bytes in, typed segments out (each with a locator: page, cell range, slide, paragraph,
lines). No file is executed: no macros, no JavaScript, no external links are fetched; spreadsheet values
are the ones saved in the file (never claimed as recalculated). A format that is only stored is said to
be so ("armazenado, mas ainda não analisável"); a failure gives a specific diagnostic instead of a
fabricated analysis. Binary content is never decoded as UTF-8 text.

The caller runs ``extract`` in a child process with a deadline (``security.broker.executor.run_in_process``)
because it parses untrusted files. That is resource containment, not a sandbox (the Execution Box is).
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field

EXTRACTOR_VERSION = "atlas-extract-1"
MAX_SEGMENT_CHARS = 4_000
MAX_CSV_ROWS_PER_SEGMENT = 60

READY = "READY_FOR_ANALYSIS"
PARTIAL = "PARTIAL"
UNSUPPORTED = "UNSUPPORTED"
FAILED = "FAILED"


@dataclass(frozen=True)
class Segment:
    locator: str  # "página 3", "Planilha1!A1:D20", "slide 2", "parágrafos 4-9", "linhas 1-60"
    kind: str  # text | table | notes | header | footer | cells
    text: str


@dataclass
class Extraction:
    state: str
    extractor: str
    segments: list[Segment] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    diagnostic: str | None = None
    missing: list[str] = field(default_factory=list)  # areas NOT extracted, e.g. "página 2" (R5-06)
    units_total: int | None = None  # pages/sheets/slides the document has, when known


def _split(locator: str, kind: str, text: str) -> list[Segment]:
    """Cut long text at paragraph/line boundaries into segments that fit the read contract."""
    text = text.strip()
    if not text:
        return []
    if len(text) <= MAX_SEGMENT_CHARS:
        return [Segment(locator, kind, text)]
    parts: list[Segment] = []
    buf = ""
    n = 1
    for line in text.splitlines(keepends=True):
        while len(line) > MAX_SEGMENT_CHARS:  # one enormous line: hard cut, still complete overall
            if buf:
                parts.append(Segment(f"{locator} (parte {n})", kind, buf.strip()))
                n += 1
                buf = ""
            parts.append(Segment(f"{locator} (parte {n})", kind, line[:MAX_SEGMENT_CHARS]))
            n += 1
            line = line[MAX_SEGMENT_CHARS:]
        if len(buf) + len(line) > MAX_SEGMENT_CHARS and buf:
            parts.append(Segment(f"{locator} (parte {n})", kind, buf.strip()))
            n += 1
            buf = ""
        buf += line
    if buf.strip():
        parts.append(Segment(f"{locator} (parte {n})", kind, buf.strip()))
    return parts


def _decode(data: bytes) -> tuple[str, list[str]]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8"), []
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), ["codificação detectada: UTF-16"]
    try:
        return data.decode("utf-8"), []
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="strict"), ["codificação detectada: Windows-1252 (não era UTF-8)"]


def _text(data: bytes, what: str) -> Extraction:
    text, warnings = _decode(data)
    if "\x00" in text:
        return Extraction(FAILED, "text", diagnostic="conteúdo binário; não é texto")
    lines = text.splitlines()
    segs: list[Segment] = []
    start, buf = 1, list[str]()
    size = 0
    for i, line in enumerate(lines, start=1):
        if size + len(line) > MAX_SEGMENT_CHARS and buf:
            segs += _split(f"linhas {start}-{i - 1}", "text", "\n".join(buf))
            start, buf, size = i, [], 0
        buf.append(line)
        size += len(line) + 1
    if buf:
        segs += _split(f"linhas {start}-{len(lines)}", "text", "\n".join(buf))
    if what == "html":
        warnings.append("HTML lido como texto-fonte; nada foi executado nem carregado da internet")
    return Extraction(READY, what, segs, warnings)


def _csv(data: bytes) -> Extraction:
    text, warnings = _decode(data)
    sample = text[:20_000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") > sample.count(",") else ","
        warnings.append(f"delimitador presumido: {delim!r}")
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    widths = {len(r) for r in rows if r}
    if len(widths) > 1:
        warnings.append(f"linhas com números de colunas diferentes: {sorted(widths)}")
    warnings.append("valores mantidos como no arquivo (sem conversão de moeda ou decimal)")
    segs: list[Segment] = []
    for i in range(0, len(rows), MAX_CSV_ROWS_PER_SEGMENT):
        chunk = rows[i : i + MAX_CSV_ROWS_PER_SEGMENT]
        body = "\n".join(delim.join(r) for r in chunk)
        segs += _split(f"linhas {i + 1}-{i + len(chunk)}", "table", body)
    return Extraction(READY, f"csv(delimiter={delim!r})", segs, warnings)


def _json(data: bytes) -> Extraction:
    text, warnings = _decode(data)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        return Extraction(FAILED, "json", diagnostic=f"JSON inválido: linha {exc.lineno}, coluna {exc.colno}")
    pretty = json.dumps(parsed, ensure_ascii=False, indent=2)
    return Extraction(READY, "json", _split("json", "text", pretty), warnings)


def _pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            return Extraction(
                FAILED, "pypdf", diagnostic="PDF protegido por senha; envie uma versão sem senha"
            )
        pages = list(reader.pages)
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        return Extraction(FAILED, "pypdf", diagnostic=f"PDF corrompido ou ilegível ({type(exc).__name__})")
    segs: list[Segment] = []
    empty: list[int] = []
    for n, page in enumerate(pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # a broken page is reported, never guessed
            empty.append(n)
            segs.append(Segment(f"página {n}", "text", f"[página ilegível: {type(exc).__name__}]"))
            continue
        if not text.strip():
            empty.append(n)
            continue
        segs += _split(f"página {n}", "text", text)
    extractor = f"pypdf {__import__('pypdf').__version__}"
    if not pages:
        return Extraction(FAILED, extractor, diagnostic="PDF sem páginas")
    if len(empty) == len(pages):
        return Extraction(
            UNSUPPORTED,
            extractor,
            diagnostic="PDF digitalizado (páginas sem texto): armazenado, mas ainda não analisável; "
            "leitura por OCR/visão ainda não está disponível",
        )
    warnings = []
    if empty:
        warnings.append(f"páginas sem texto extraível (imagem/digitalizadas): {empty}")
    return Extraction(
        PARTIAL if empty else READY,
        extractor,
        segs,
        warnings,
        missing=[f"página {n}" for n in empty],
        units_total=len(pages),
    )


def _docx(data: bytes) -> Extraction:
    import docx
    from docx.oxml.ns import qn

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        return Extraction(
            FAILED, "python-docx", diagnostic=f"DOCX corrompido ou ilegível ({type(exc).__name__})"
        )
    segs: list[Segment] = []
    para_buf: list[str] = []
    para_start = 1
    p_no = t_no = 0

    def flush(end: int) -> None:
        nonlocal para_buf, para_start
        if para_buf:
            segs.extend(_split(f"parágrafos {para_start}-{end}", "text", "\n".join(para_buf)))
        para_buf = []
        para_start = end + 1

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p_no += 1
            text = "".join(t.text or "" for t in child.iter(qn("w:t")))
            if text.strip():
                para_buf.append(text)
            if sum(len(x) for x in para_buf) > MAX_SEGMENT_CHARS:
                flush(p_no)
        elif child.tag == qn("w:tbl"):
            flush(p_no)
            t_no += 1
            rows = []
            for tr in child.iter(qn("w:tr")):
                cells = ["".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in tr.iter(qn("w:tc"))]
                rows.append(" | ".join(cells))
            segs += _split(f"tabela {t_no}", "table", "\n".join(rows))
    flush(p_no)
    seen: set[str] = set()
    for s_no, section in enumerate(document.sections, start=1):
        for kind, part in (("header", section.header), ("footer", section.footer)):
            text = "\n".join(p.text for p in part.paragraphs if p.text.strip())
            if text and text not in seen:
                seen.add(text)
                segs += _split(f"{'cabeçalho' if kind == 'header' else 'rodapé'} (seção {s_no})", kind, text)
    warnings = ["revisões controladas: lido o texto aparente; comentários não incluídos"]
    return Extraction(
        READY,
        f"python-docx {docx.__version__ if hasattr(docx, '__version__') else ''}".strip(),
        segs,
        warnings,
    )


def _xlsx(data: bytes) -> Extraction:
    import openpyxl
    from openpyxl.utils import get_column_letter

    try:
        formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        return Extraction(
            FAILED, "openpyxl", diagnostic=f"XLSX corrompido ou ilegível ({type(exc).__name__})"
        )
    warnings = ["valores de fórmulas: os salvos no arquivo (não recalculados)"]
    if getattr(formulas, "_external_links", None):
        warnings.append("a planilha tem vínculos externos; não foram abertos nem atualizados")
    segs: list[Segment] = []
    for ws in formulas.worksheets:
        vs = values[ws.title]
        hidden_sheet = ws.sheet_state != "visible"
        if hidden_sheet:
            warnings.append(f"aba oculta incluída: {ws.title}")
        hidden_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
        hidden_cols = [c for c, d in ws.column_dimensions.items() if d.hidden]
        if hidden_rows or hidden_cols:
            warnings.append(
                f"{ws.title}: linhas ocultas {hidden_rows[:10]} colunas ocultas {hidden_cols[:10]}"
            )
        lines: list[str] = []
        first_row = None
        last_row = 0
        max_col = 1
        for row in ws.iter_rows():
            cells = []
            for cell in row:
                if cell.value is None:
                    continue
                ref = f"{get_column_letter(cell.column)}{cell.row}"
                raw = cell.value
                saved = vs[ref].value
                if isinstance(raw, str) and raw.startswith("="):
                    cells.append(f"{ref}={saved!r} (fórmula {raw})")
                else:
                    cells.append(f"{ref}={raw!r}")
                max_col = max(max_col, int(cell.column or 1))
            if cells:
                first_row = first_row or row[0].row
                last_row = int(row[0].row or last_row)
                lines.append("; ".join(cells))
            if sum(len(x) for x in lines) > MAX_SEGMENT_CHARS - 500:
                segs += _split(
                    f"{ws.title}!A{first_row}:{get_column_letter(max_col)}{last_row}",
                    "cells",
                    "\n".join(lines),
                )
                lines, first_row = [], None
        if lines:
            segs += _split(
                f"{ws.title}!A{first_row}:{get_column_letter(max_col)}{last_row}", "cells", "\n".join(lines)
            )
    return Extraction(READY, f"openpyxl {openpyxl.__version__}", segs, warnings)


def _pptx(data: bytes) -> Extraction:
    import pptx
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        return Extraction(
            FAILED, "python-pptx", diagnostic=f"PPTX corrompido ou ilegível ({type(exc).__name__})"
        )
    segs: list[Segment] = []
    pictures = 0
    for n, slide in enumerate(deck.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [" | ".join(c.text for c in r.cells) for r in shape.table.rows]
                segs += _split(f"slide {n} (tabela)", "table", "\n".join(rows))
            if shape.shape_type == 13:  # picture
                pictures += 1
        segs += _split(f"slide {n}", "text", "\n".join(texts))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            segs += _split(f"slide {n} (notas)", "notes", slide.notes_slide.notes_text_frame.text)
    warnings = [f"{pictures} imagem(ns) não interpretada(s) (visão ainda não disponível)"] if pictures else []
    return Extraction(READY, f"python-pptx {pptx.__version__}", segs, warnings)


def extract(data: bytes, name: str) -> Extraction:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext in ("txt", "md", "html"):
        return _text(data, "html" if ext == "html" else "text")
    if ext == "csv":
        return _csv(data)
    if ext == "json":
        return _json(data)
    if ext == "pdf":
        return _pdf(data)
    if ext == "docx":
        return _docx(data)
    if ext == "xlsx":
        return _xlsx(data)
    if ext == "pptx":
        return _pptx(data)
    if ext in ("png", "jpg", "jpeg"):
        return Extraction(
            UNSUPPORTED,
            "none",
            diagnostic="imagem armazenada, mas ainda não analisável: "
            "leitura visual/OCR ainda não está disponível",
        )
    return Extraction(UNSUPPORTED, "none", diagnostic=f"formato .{ext} armazenado, mas ainda não analisável")


def extract_job(job: dict[str, object]) -> Extraction:
    """Entry point for the child process (picklable arguments, picklable result)."""
    data = job["data"]
    name = job["name"]
    assert isinstance(data, bytes) and isinstance(name, str)
    return extract(data, name)
