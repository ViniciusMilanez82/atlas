"""Deliverable generation in real formats (N11, spec 14.3).

The model proposes structured content (blocks); trusted code renders it into PDF, DOCX, XLSX or PPTX and
then VALIDATES the file by re-reading it with the same extractors used for inputs: every text block must
be found again, tables must keep their cells, and the file must open. A file that does not pass is never
stored as a deliverable. Spreadsheet formulas are preserved but not claimed as recalculated.

Blocks: {"type": "heading"|"paragraph"|"bullets"|"table"|"slide", ...}
  heading:   {"text": str}
  paragraph: {"text": str}
  bullets:   {"items": [str]}
  table:     {"rows": [[str|number]], "sheet": str?}   (first row = header)
  slide:     {"title": str, "bullets": [str], "notes": str?}   (PPTX; other formats render it as a section)
"""

from __future__ import annotations

import io
import textwrap
from dataclasses import dataclass
from typing import Any

from runtime.documents.extract import READY, extract

FORMATS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
MAX_BLOCKS = 400


class GenerationError(ValueError):
    pass


@dataclass(frozen=True)
class Generated:
    data: bytes
    mime: str
    pages_or_parts: int
    warnings: list[str]


def _cell(v: Any) -> str:
    return "" if v is None else (f"{v}" if not isinstance(v, float) else f"{v:g}")


def _texts(blocks: list[dict[str, Any]]) -> list[str]:
    """Every piece of text that must survive the round trip."""
    out: list[str] = []
    for b in blocks:
        t = b.get("type")
        if t in ("heading", "paragraph"):
            out.append(str(b.get("text", "")))
        elif t == "bullets":
            out += [str(i) for i in b.get("items", [])]
        elif t == "table":
            out += [
                _cell(c)
                for row in b.get("rows", [])
                for c in row
                if not (isinstance(c, str) and c.startswith("="))
            ]
        elif t == "slide":
            out.append(str(b.get("title", "")))
            out += [str(i) for i in b.get("bullets", [])]
    return [x for x in out if x.strip()]


# ---------------------------------------------------------------- PDF (standard font, WinAnsi: Portuguese OK)


def _pdf(title: str, blocks: list[dict[str, Any]]) -> tuple[bytes, int]:
    lines: list[tuple[str, int]] = []  # (text, font size)

    def add(text: str, size: int, width: int) -> None:
        for para in str(text).splitlines() or [""]:
            wrapped = textwrap.wrap(para, width=width) or [""]
            lines.extend((w, size) for w in wrapped)

    add(title, 16, 55)
    lines.append(("", 11))
    for b in blocks:
        t = b.get("type")
        if t == "heading":
            lines.append(("", 11))
            add(b.get("text", ""), 13, 70)
        elif t == "paragraph":
            add(b.get("text", ""), 11, 90)
            lines.append(("", 11))
        elif t == "bullets":
            for item in b.get("items", []):
                add("- " + str(item), 11, 88)
        elif t == "table":
            for row in b.get("rows", []):
                add(" | ".join(_cell(c) for c in row), 10, 100)
            lines.append(("", 11))
        elif t == "slide":
            add(b.get("title", ""), 13, 70)
            for item in b.get("bullets", []):
                add("- " + str(item), 11, 88)
    pages: list[list[tuple[str, int]]] = []
    y = 750
    current: list[tuple[str, int]] = []
    for text, size in lines:
        step = size + 4
        if y - step < 60:
            pages.append(current)
            current, y = [], 750
        current.append((text, size))
        y -= step
    pages.append(current)

    def esc(s: str) -> bytes:
        raw = s.encode("cp1252", errors="replace")
        return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")

    objs: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    font = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    objs.append(font)
    kids = []
    for page in pages:
        ops = [b"BT"]
        y = 750
        for text, size in page:
            ops.append(b"/F1 %d Tf 1 0 0 1 60 %d Tm (" % (size, y) + esc(text) + b") Tj")
            y -= size + 4
        ops.append(b"ET")
        stream = b"\n".join(ops)
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        content_no = len(objs)
        objs.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R"
            b" /Resources << /Font << /F1 3 0 R >> >> >>" % content_no
        )
        kids.append(len(objs))
    objs[1] = (
        b"<< /Type /Pages /Kids [" + b" ".join(b"%d 0 R" % k for k in kids) + b"] /Count %d >>" % len(kids)
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue(), len(pages)


def _docx(title: str, blocks: list[dict[str, Any]]) -> tuple[bytes, int]:
    import docx

    d = docx.Document()
    d.add_heading(title, level=0)
    for b in blocks:
        t = b.get("type")
        if t == "heading":
            d.add_heading(str(b.get("text", "")), level=1)
        elif t == "paragraph":
            d.add_paragraph(str(b.get("text", "")))
        elif t == "bullets":
            for item in b.get("items", []):
                d.add_paragraph(str(item), style="List Bullet")
        elif t == "table":
            rows = b.get("rows", [])
            if rows:
                table = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for i, row in enumerate(rows):
                    for j, c in enumerate(row):
                        table.cell(i, j).text = _cell(c)
        elif t == "slide":
            d.add_heading(str(b.get("title", "")), level=1)
            for item in b.get("bullets", []):
                d.add_paragraph(str(item), style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue(), 1


def _xlsx(title: str, blocks: list[dict[str, Any]]) -> tuple[bytes, int]:
    import openpyxl

    wb = openpyxl.Workbook()
    first = True
    n = 0
    notes: list[str] = []
    for b in blocks:
        if b.get("type") != "table":
            if b.get("type") in ("heading", "paragraph"):
                notes.append(str(b.get("text", "")))
            elif b.get("type") == "bullets":
                notes += [str(i) for i in b.get("items", [])]
            elif b.get("type") == "slide":
                notes += [str(b.get("title", ""))] + [str(i) for i in b.get("bullets", [])]
            continue
        n += 1
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        first = False
        ws.title = str(b.get("sheet") or f"Tabela {n}")[:31]
        for row in b.get("rows", []):
            ws.append(
                [
                    c
                    if isinstance(c, (int, float)) or (isinstance(c, str) and c.startswith("="))
                    else _cell(c)
                    for c in row
                ]
            )
    if notes:
        ws = wb.active if first else wb.create_sheet("Notas")
        assert ws is not None
        if first:
            ws.title = "Notas"
        ws.append([title])
        for line in notes:
            ws.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(wb.worksheets)


def _pptx(title: str, blocks: list[dict[str, Any]]) -> tuple[bytes, int]:
    from pptx import Presentation

    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    for b in blocks:
        t = b.get("type")
        if t == "slide":
            s = deck.slides.add_slide(deck.slide_layouts[1])
            s.shapes.title.text = str(b.get("title", ""))
            s.placeholders[1].text = "\n".join(str(i) for i in b.get("bullets", []))
            if b.get("notes"):
                s.notes_slide.notes_text_frame.text = str(b["notes"])
        elif t in ("heading", "paragraph"):
            s = deck.slides.add_slide(deck.slide_layouts[1])
            s.shapes.title.text = str(b.get("text", ""))[:120] if t == "heading" else title
            if t == "paragraph":
                s.placeholders[1].text = str(b.get("text", ""))
        elif t == "bullets":
            s = deck.slides.add_slide(deck.slide_layouts[1])
            s.shapes.title.text = title
            s.placeholders[1].text = "\n".join(str(i) for i in b.get("items", []))
        elif t == "table":
            rows = b.get("rows", [])
            if rows:
                s = deck.slides.add_slide(deck.slide_layouts[5])
                s.shapes.title.text = str(b.get("sheet") or title)
                cols = max(len(r) for r in rows)
                shape = s.shapes.add_table(len(rows), cols, 457200, 1600000, 8229600, 400000 * len(rows))
                for i, row in enumerate(rows):
                    for j, c in enumerate(row):
                        shape.table.cell(i, j).text = _cell(c)
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue(), len(deck.slides)


def generate(name: str, title: str, blocks: list[dict[str, Any]]) -> Generated:
    ext = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in FORMATS:
        raise GenerationError(
            f"formato {ext or '(nenhum)'} não é gerado aqui; use .pdf, .docx, .xlsx ou .pptx"
        )
    if not blocks or len(blocks) > MAX_BLOCKS:
        raise GenerationError("o documento precisa de 1 a 400 blocos")
    maker = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".pptx": _pptx}[ext]
    data, parts = maker(title, blocks)
    # Validation by round trip with the input extractors: it opens and every text came back.
    ex = extract(data, name)
    if ex.state != READY:
        raise GenerationError(f"o arquivo gerado não passou na releitura: {ex.diagnostic or ex.state}")
    extracted = " ".join(s.text for s in ex.segments)
    squashed = " ".join(extracted.split())
    missing = [t for t in _texts(blocks) if " ".join(t.split()) not in squashed]
    if missing:
        raise GenerationError(f"textos ausentes na releitura do arquivo gerado: {missing[:3]}")
    warnings = ["fórmulas preservadas; valores não recalculados aqui"] if ext == ".xlsx" else []
    return Generated(data, FORMATS[ext], parts, warnings + ex.warnings[:3])
