"""A3-09 (receiving a format is not understanding it) and A3-10 (long content is never cut silently) -
scenarios T12 and T13.

The same synthetic content in several real formats (built here with the writer side of the libraries, or
byte by byte for PDF) is imported through the real IPC path and read with the paths the agent uses.
Expected values ("gabarito") are fixed before running.
"""

from __future__ import annotations

import base64
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import pytest

from runtime.artifacts.manager import ArtifactManager
from runtime.documents.store import DocumentStore
from shared.ids import new_id
from tests.conftest import World
from tests.integration.test_agent_loop import SPEC, UUID, build, decision
from tests.integration.test_alpha2 import ok

VALUE = "12.345,67"
CLAUSE = "Clausula final: multa de 7% por atraso"


def pdf_bytes(pages: list[str]) -> bytes:
    """Minimal, valid PDF with one Helvetica text block per page (empty string = page without text)."""
    objs: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    kids = []
    for text in pages:
        content = (
            b"BT /F1 12 Tf 72 720 Td ("
            + text.replace("(", "[").replace(")", "]").encode("latin-1")
            + b") Tj ET"
            if text
            else b""
        )
        objs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
        content_no = len(objs)
        objs.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R"
            b" /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> >>"
            % content_no
        )
        kids.append(len(objs))
    objs[1] = (
        b"<< /Type /Pages /Kids [" + b" ".join(b"%d 0 R" % k for k in kids) + b"] /Count %d >>" % len(kids)
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Contrato sintetico de fornecimento.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Valor"
    t.cell(1, 0).text, t.cell(1, 1).text = "Total", VALUE
    d.add_paragraph(CLAUSE)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Resumo"
    ws["A1"], ws["B1"] = "Item", "Valor"
    ws["A2"], ws["B2"] = "Total", VALUE
    ws["A3"] = CLAUSE
    hidden = wb.create_sheet("Oculta")
    hidden["A1"] = "dado de aba oculta"
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pptx_bytes() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    s = deck.slides.add_slide(deck.slide_layouts[1])
    s.shapes.title.text = "Contrato sintetico"
    s.placeholders[1].text = f"Total {VALUE}"
    s2 = deck.slides.add_slide(deck.slide_layouts[1])
    s2.shapes.title.text = "Condicoes"
    s2.placeholders[1].text = CLAUSE
    s2.notes_slide.notes_text_frame.text = "nota do apresentador"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _upload(c: Any, world: World, name: str, data: bytes) -> dict[str, Any]:
    ref = new_id()
    ok(c.call("artifacts.upload", upload_ref=ref, offset=0, data_b64=base64.b64encode(data).decode()))
    return ok(c.call("artifacts.import", employee_id=world.employee.id, upload_ref=ref, declared_name=name))


def _all_text(world: World, env: Any, artifact_id: str) -> str:
    ds = DocumentStore(world.conn, world.clock, ArtifactManager(world.conn, world.clock, env.store_root))
    out, cursor = [], 0
    while True:
        page = ds.read(artifact_id, task_id=None, cursor=cursor)
        out += [f"[{s['locator']}] {s['text']}" for s in page["segments"]]
        if not page["has_more"]:
            return "\n".join(out)
        cursor = page["next_cursor"]


@pytest.mark.parametrize(
    ("name", "maker", "state", "locator"),
    [
        (
            "contrato.txt",
            lambda: f"Contrato sintetico\nTotal: {VALUE}\n{CLAUSE}\n".encode(),
            "READY_FOR_ANALYSIS",
            "linhas",
        ),
        (
            "contrato.csv",
            lambda: f"item;valor\nTotal;{VALUE}\nObs;{CLAUSE}\n".encode(),
            "READY_FOR_ANALYSIS",
            "linhas",
        ),
        ("contrato.pdf", lambda: pdf_bytes([f"Total {VALUE}", CLAUSE]), "READY_FOR_ANALYSIS", "página 2"),
        ("contrato.docx", docx_bytes, "READY_FOR_ANALYSIS", "tabela 1"),
        ("contrato.xlsx", xlsx_bytes, "READY_FOR_ANALYSIS", "Resumo!A"),
        ("contrato.pptx", pptx_bytes, "READY_FOR_ANALYSIS", "slide 2"),
    ],
)
def test_each_format_is_understood_with_provenance(
    env: Any, world: World, name: str, maker: Any, state: str, locator: str
) -> None:
    c = env.connect()
    imported = _upload(c, world, name, maker())
    assert imported["analysis"]["state"] == state, imported["analysis"]
    text = _all_text(world, env, imported["artifact_id"])
    assert VALUE in text and CLAUSE in text  # gabarito: table value and the final clause
    assert locator in text  # page / cell range / slide / table reference
    assert "�" not in text  # before the fix: binary formats were decoded as UTF-8 with replacement


def test_csv_keeps_raw_numbers_and_xlsx_reports_hidden_sheets(env: Any, world: World) -> None:
    c = env.connect()
    csv_art = _upload(c, world, "valores.csv", f"a;b\nx;{VALUE}\n".encode())
    assert VALUE in _all_text(world, env, csv_art["artifact_id"])  # not converted to 12345.67
    xl = _upload(c, world, "p.xlsx", xlsx_bytes())
    assert any("oculta" in w for w in xl["analysis"]["warnings"])
    assert any("não recalculados" in w for w in xl["analysis"]["warnings"])


def test_scanned_pdf_image_and_corrupted_file_are_honest(env: Any, world: World) -> None:
    c = env.connect()
    scanned = _upload(c, world, "digitalizado.pdf", pdf_bytes(["", ""]))
    assert scanned["analysis"]["state"] == "UNSUPPORTED"
    assert "ainda não analisável" in scanned["analysis"]["diagnostic"]
    img = _upload(c, world, "foto.png", PNG)
    assert img["analysis"]["state"] == "UNSUPPORTED"
    broken = io.BytesIO()
    with zipfile.ZipFile(broken, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", "isto nao e xml valido <<<")
    bad = _upload(c, world, "quebrado.docx", broken.getvalue())
    assert bad["analysis"]["state"] == "FAILED" and bad["analysis"]["diagnostic"]


def test_decisive_clause_after_20k_and_200k_is_reached_by_paged_reading(world: World, tmp_path: Path) -> None:
    filler = "Linha de texto sintetico sem importancia para a decisao. " * 3
    lines = (
        [filler] * 400
        + ["CLAUSULA-A: prazo maximo de 45 dias."]
        + [filler] * 3200
        + ["CLAUSULA-B: rescisao sem multa apos 90 dias."]
    )
    doc = "\n".join(lines)
    assert doc.index("CLAUSULA-A") > 20_000 and doc.index("CLAUSULA-B") > 200_000
    seen: dict[str, Any] = {"observations": []}

    def policy(turn: int, prompt: str) -> str:
        found = re.findall(r'"next_cursor": (\d+)', prompt)
        if turn == 1:
            ids = re.findall(rf"artifact_id ({UUID})", prompt)
            seen["id"] = ids[0]
            return decision("tool", tool="documents.read", inp={"artifact_id": ids[0], "cursor": 0})
        if '"has_more": true' in prompt.split("EXTERNAL-DATA")[-1] or (found and "CLAUSULA-B" not in prompt):
            return decision(
                "tool", tool="documents.read", inp={"artifact_id": seen["id"], "cursor": int(found[-1])}
            )
        return decision("ask_owner", question="Li o documento inteiro. Posso resumir?")

    s = build(world, tmp_path, policy)
    tid = s.tasks.create(
        world.owner,
        employee_id=world.employee.id,
        objective="Revise o contrato inteiro",
        criteria=[("x", True)],
    )
    f = tmp_path / "longo.txt"
    f.write_text(doc, encoding="utf-8")
    s.am.import_file(f, actor=world.owner, employee_id=world.employee.id, task_id=tid)
    out = s.runner.run(tid, SPEC)
    prompts = "\n".join(s.model.prompts)
    assert (
        "CLAUSULA-A" in prompts and "CLAUSULA-B" in prompts
    )  # before the fix: text[:200_000], JSON[:20_000]
    assert out.state == "WAITING_USER"
    for (content,) in world.conn.execute("SELECT content FROM step_observations WHERE task_id = ?", (tid,)):
        json.loads(content)  # every observation is valid JSON: nothing cut in the middle
    art = world.conn.execute("SELECT artifact_id FROM artifact_links WHERE task_id = ?", (tid,)).fetchone()[0]
    cov = DocumentStore(world.conn, world.clock, s.am).coverage(tid, art)
    assert cov["complete"] is True and cov["read_segments"] == cov["total_segments"]
